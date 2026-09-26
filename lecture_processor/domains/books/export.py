"""A4 imposition shared by faithful DOCX, editable DOCX and PDF exports."""
from __future__ import annotations

import base64
import copy
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsmap, qn
from docx.shared import Mm, Pt
from PIL import Image
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader

from . import model

nsmap.setdefault('v', 'urn:schemas-microsoft-com:vml')
nsmap.setdefault('o', 'urn:schemas-microsoft-com:office:office')
PT = 72 / 25.4


def decode_image(value):
    if not isinstance(value, str) or not value.startswith('data:image/png;base64,') or len(value) > 20_000_000:
        raise model.BookError('A page preview is missing or too large. Please try exporting again.')
    try:
        data = base64.b64decode(value.split(',', 1)[1], validate=True)
        image = Image.open(io.BytesIO(data))
        if image.format != 'PNG' or image.width * image.height > 20_000_000:
            raise ValueError('Invalid page image')
        image.load()
        return image.convert('RGBA')
    except (ValueError, OSError) as error:
        raise model.BookError('A page preview could not be read. Please retry the export.') from error


def anchor_picture(paragraph, image, x, y, w, h):
    out = io.BytesIO()
    image.save(out, format='PNG')
    out.seek(0)
    run = paragraph.add_run()
    inline = run.add_picture(out, width=Mm(w), height=Mm(h))._inline
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':'0', 'distB':'0', 'distL':'0', 'distR':'0', 'simplePos':'0', 'relativeHeight':'0', 'behindDoc':'1', 'locked':'0', 'layoutInCell':'1', 'allowOverlap':'1'}.items():
        anchor.set(key, value)
    simple = OxmlElement('wp:simplePos')
    simple.set('x', '0'); simple.set('y', '0')
    anchor.append(simple)
    for axis, value in [('H', x), ('V', y)]:
        pos = OxmlElement('wp:position' + axis)
        pos.set('relativeFrom', 'page')
        offset = OxmlElement('wp:posOffset')
        offset.text = str(int(Mm(value)))
        pos.append(offset); anchor.append(pos)
    for tag in ('wp:extent', 'wp:effectExtent'):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def native_object(paragraph, obj, side, ordinal):
    x = (obj['x'] + side * model.WIDTH) * PT
    y = obj['y'] * PT
    width, height = obj['w'] * PT, obj['h'] * PT
    tag = 'v:oval' if obj.get('shape') == 'circle' else 'v:rect'
    if obj['type'] == 'text':
        tag = 'v:rect'
    pict = OxmlElement('w:pict')
    shape = OxmlElement(tag)
    shape.set('id', 'book-object-' + str(ordinal))
    shape.set('style', f'position:absolute;margin-left:{x:.3f}pt;margin-top:{y:.3f}pt;width:{width:.3f}pt;height:{height:.3f}pt;rotation:{obj["rotation"]};z-index:{ordinal + 1};mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
    shape.set('filled', 'f' if obj['type'] == 'text' else 't')
    shape.set('stroked', 'f' if obj['type'] == 'text' else 't')
    shape.set('fillcolor', obj['fill'])
    shape.set('strokecolor', obj['stroke'])
    shape.set('strokeweight', f'{obj["strokeWidth"] * PT:.2f}pt')
    if obj['type'] == 'text':
        textbox = OxmlElement('v:textbox')
        textbox.set('inset', '0,0,0,0')
        content = OxmlElement('w:txbxContent')
        p = OxmlElement('w:p')
        pp = OxmlElement('w:pPr')
        align = OxmlElement('w:jc'); align.set(qn('w:val'), obj['style']['align']); pp.append(align)
        spacing = OxmlElement('w:spacing'); spacing.set(qn('w:before'), '0'); spacing.set(qn('w:after'), '0'); spacing.set(qn('w:line'), str(round(obj['style']['lineHeight'] * obj['style']['size'] * 20))); spacing.set(qn('w:lineRule'), 'exact'); pp.append(spacing)
        p.append(pp)
        runs = obj.get('runs') or [{'text': obj['text'], 'style': obj['style']}]
        for source in runs:
            style = source['style']
            for idx, line in enumerate(source['text'].split('\n')):
                r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
                fonts = OxmlElement('w:rFonts')
                for name in ('ascii', 'hAnsi', 'cs'):
                    fonts.set(qn('w:' + name), style['font'])
                props.append(fonts)
                size = OxmlElement('w:sz'); size.set(qn('w:val'), str(round(style['size'] * 2))); props.append(size)
                color = OxmlElement('w:color'); color.set(qn('w:val'), style['color'][1:]); props.append(color)
                tracking = OxmlElement('w:spacing'); tracking.set(qn('w:val'), str(round(style['letterSpacing'] * 15))); props.append(tracking)
                for flag, element in [('italic', 'w:i'), ('underline', 'w:u')]:
                    if style[flag]:
                        el = OxmlElement(element)
                        if flag == 'underline': el.set(qn('w:val'), 'single')
                        props.append(el)
                if style['weight'] >= 600: props.append(OxmlElement('w:b'))
                if style['highlightOn']:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), style['highlight'][1:]); props.append(shade)
                r.append(props)
                if idx: r.append(OxmlElement('w:br'))
                t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = line; r.append(t); p.append(r)
        content.append(p); textbox.append(content); shape.append(textbox)
    pict.append(shape); paragraph.add_run()._r.append(pict)


def native_eligible(obj):
    return not obj.get('hidden') and obj['opacity'] == 1 and not obj.get('spanId') and (
        (obj['type'] == 'text' and not obj['style']['outline'] and not any(r['style']['outline'] for r in obj.get('runs', []))) or
        (obj['type'] == 'shape' and obj['shape'] in ('square', 'circle')))


def ink_saving_object(obj):
    """Keep native Word text as readable as the shared white-paper renderer."""
    obj = copy.deepcopy(obj)
    for style in [obj['style']] + [run['style'] for run in obj.get('runs', [])]:
        channels = [int(style['color'][i:i + 2], 16) / 255 for i in (1, 3, 5)]
        linear = [v / 12.92 if v <= .04045 else ((v + .055) / 1.055) ** 2.4 for v in channels]
        luminance = sum(value * factor for value, factor in zip(linear, (.2126, .7152, .0722)))
        if 1.05 / (luminance + .05) < 4.5:
            style['color'] = '#062940'
    return obj


def generate(raw):
    pages = [model.page(p) for p in model.array(raw.get('pages'), 100, 'page list')]
    model.validate_order([p['id'] for p in pages])
    previews = raw.get('previews', [])
    if len(previews) != len(pages):
        raise model.BookError('Some pages are still preparing. Please retry the export.')
    mode = raw.get('format', 'faithful')
    if mode not in ('faithful', 'editable', 'pdf'):
        raise model.BookError('Choose Word or PDF.')
    pairs = model.sheets(pages, raw.get('arrangement', 'cut'))
    output = io.BytesIO()
    economy = bool(raw.get('economy'))
    if mode == 'pdf':
        canvas = Canvas(output, pagesize=(297 * PT, 210 * PT))
        canvas.setTitle(model.text(raw.get('title'), 150) or 'Book')
        for pair in pairs:
            for side, index in enumerate(pair):
                if index is not None:
                    image = decode_image(previews[index]).convert('RGB') if not economy else decode_image(previews[index]).convert('L').convert('RGB')
                    canvas.drawImage(ImageReader(image), side * model.WIDTH * PT, 0, width=model.WIDTH * PT, height=210 * PT)
            if raw.get('guides'):
                canvas.setStrokeColorRGB(.75, .75, .75); canvas.setDash(2, 3)
                canvas.line(model.WIDTH * PT, 0, model.WIDTH * PT, 210 * PT)
            canvas.showPage()
        canvas.save()
        return output.getvalue(), 'application/pdf', 'pdf'
    doc = Document()
    doc.core_properties.title = model.text(raw.get('title'), 150) or 'Book'
    doc.core_properties.author = ''
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(0)
    normal = doc.styles['Normal']; normal.font.size = Pt(1)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = Pt(1)
    ordinal = 0
    for sheet, pair in enumerate(pairs):
        paragraph = doc.add_paragraph()
        if sheet:
            paragraph.paragraph_format.page_break_before = True
        for side, index in enumerate(pair):
            if index is None:
                continue
            img = decode_image(previews[index])
            if economy: img = img.convert('L').convert('RGBA')
            anchor_picture(paragraph, img, side * model.WIDTH, 0, model.WIDTH, 210)
            if mode == 'editable':
                for obj in pages[index]['items']:
                    if native_eligible(obj):
                        ordinal += 1
                        if economy and (pages[index].get('decoration') or {}).get('id') == 'xped':
                            obj = ink_saving_object(obj)
                        native_object(paragraph, obj, side, ordinal)
        if raw.get('guides'):
            pict = OxmlElement('w:pict')
            line = parse_xml('<v:line xmlns:v="urn:schemas-microsoft-com:vml" style="position:absolute;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:99999" from="420.945pt,0pt" to="420.945pt,595.276pt" strokecolor="#bbbbbb" strokeweight="0.25pt"/>')
            pict.append(line); paragraph.add_run()._r.append(pict)
    doc.save(output)
    return output.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'docx'
